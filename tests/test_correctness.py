"""
Regression tests for correctness fixes across scores, filters, groups, exports, derive, signals, and presets.
All tests run offline on synthetic data.
"""
import datetime as dt
import io
import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from screener import export
from screener import groups as G
from screener import i18n as I
from screener import presets as P
from screener import signals as S
from screener.derive import add_derived
from screener.fields import COLUMNS, CANDLES
from screener.filters import apply_filters, option_labels, FILTERS_BY_KEY
from screener.scores import add_scores


def frame_correctness(**cols) -> pd.DataFrame:
    """Minimal raw frame with synthetic data."""
    n = len(next(iter(cols.values())))
    base = {k: [np.nan] * n for k, c in COLUMNS.items() if c.source == 'provider' and c.kind != 'text'}
    base.update({
        "ticker": [f"T{i}" for i in range(n)],
        "company": [f"Co {i}" for i in range(n)],
        "sector": ["Tech"] * n,
        "industry": ["Software"] * n,
        "country": ["United States"] * n,
        "exchange": ["NASDAQ"] * n,
        "indexes": [None] * n,
        "earnings_date": [pd.NaT] * n,
        "tv_symbol": [f"NASDAQ:T{i}" for i in range(n)],
    })
    for k in CANDLES:
        base[k] = [0.0] * n
    base.update(cols)
    return pd.DataFrame(base)


def test_quality_score_debt_eq_ranking():
    # Repro: debt_eq = [0.0, 0.5, 3.0, -2.0]
    # 0.0 is best (no debt), 0.5 is low debt, 3.0 is high debt, -2.0 is negative equity (worst).
    raw = frame_correctness(price=[10.0] * 4, debt_eq=[0.0, 0.5, 3.0, -2.0],
                            roe=[15.0] * 4, roic=[10.0] * 4, gross_margin=[40.0] * 4)
    s = add_scores(raw)
    # Quality score for 0.0 debt must rank best
    assert s.loc[0, "score_quality"] == s["score_quality"].max()
    assert s.loc[0, "score_quality"] > s.loc[1, "score_quality"] > s.loc[2, "score_quality"] > s.loc[3, "score_quality"]
    assert s.loc[3, "score_quality"] == s["score_quality"].min()


def test_earnings_date_not_in_next_days():
    # Missing earnings date or past earnings date (NaN days_to_earnings) must not pass
    raw = frame_correctness(price=[10.0] * 3)
    d = add_derived(raw, today=dt.date(2026, 10, 20))
    d.loc[0, "days_to_earnings"] = 10.0   # in 10 days
    d.loc[1, "days_to_earnings"] = 3.0    # in 3 days
    d.loc[2, "days_to_earnings"] = np.nan # no earnings date / past

    pass7 = apply_filters(d, {"earnings": "Not in next 7 days"})
    assert pass7["ticker"].tolist() == ["T0"]

    pass14 = apply_filters(d, {"earnings": "Not in next 14 days"})
    assert pass14["ticker"].tolist() == []


def test_under_x_and_low_filter_boundaries():
    # LT Debt/Eq "Under 0.1" / "Under 1" must exclude negative equity (-5.0)
    # Low (<0.1) must exclude exact boundary 0.1
    d = frame_correctness(price=[10.0] * 4, lt_debt_eq=[0.05, 0.1, 0.5, -5.0],
                          debt_eq=[0.05, 0.1, 0.5, -5.0], payout=[10.0, 40.0, 80.0, -10.0])

    pass_lt_under1 = apply_filters(d, {"lt_debt_eq": "Under 1"})
    assert pass_lt_under1["ticker"].tolist() == ["T0", "T1", "T2"]

    pass_lt_low = apply_filters(d, {"lt_debt_eq": "Low (<0.1)"})
    assert pass_lt_low["ticker"].tolist() == ["T0"]

    pass_payout_under40 = apply_filters(d, {"payout": "Under 40%"})
    assert pass_payout_under40["ticker"].tolist() == ["T0"]


def test_recom_hold_boundary():
    # Recom 2.5 is Buy or better; Hold is > 2.5 and <= 3.5
    d = frame_correctness(price=[10.0] * 3, recom=[2.5, 2.6, 3.5])

    pass_buy = apply_filters(d, {"recom": "Buy or better"})
    assert "T0" in pass_buy["ticker"].tolist()

    pass_hold = apply_filters(d, {"recom": "Hold"})
    assert pass_hold["ticker"].tolist() == ["T1", "T2"]


def test_composite_score_reranked():
    # Check that composite score is re-ranked over 0-100 percentile
    d = frame_correctness(
        price=[10.0] * 100,
        pe=np.linspace(10, 50, 100), ps=np.linspace(1, 10, 100), pb=np.linspace(1, 5, 100),
        roe=np.linspace(5, 25, 100), roic=np.linspace(5, 20, 100), gross_margin=np.linspace(20, 60, 100),
        perf_q=np.linspace(-10, 30, 100), perf_h=np.linspace(-10, 50, 100), perf_y=np.linspace(-10, 70, 100),
    )
    s = add_scores(d)
    assert s["score"].max() == 100
    assert s["score"].min() == 1
    # Over 90 should match approximately the top 10%
    top90 = s[s["score"] > 90]
    assert 8 <= len(top90) <= 12


def test_group_median_averages_and_guide():
    d = frame_correctness(price=[10.0] * 5, change=[2.0, 3.0, 2.5, 1.0, 2500.0],
                          sector=["Tech"] * 5, market_cap=[1e9] * 5)
    grp_med = G.group_performance(d, by="sector", weight="median")
    assert grp_med.loc[0, "change"] == pytest.approx(2.5)

    grp_eq = G.group_performance(d, by="sector", weight="equal")
    assert grp_eq.loc[0, "change"] > 500.0


def test_derive_uses_ny_market_date():
    raw = frame_correctness(price=[10.0], earnings_date=[pd.Timestamp("2026-10-21")])
    # Fetch time is 01:00 Israel time on Oct 21 -> NY date is Oct 20
    ny_date = dt.date(2026, 10, 20)
    d = add_derived(raw, today=ny_date)
    assert d.loc[0, "days_to_earnings"] == 1.0   # Tomorrow relative to NY Oct 20


def test_export_formula_injection_and_illegal_chars():
    d = frame_correctness(
        price=[10.0], company=["=CMD|' /C calc'!A0"], ticker=["+AAPL"],
        sector=["-Tech"], industry=["@Software"], range52=[0.49]
    )
    d.loc[0, "company"] = "Bad\x0bCompany"

    excel_bytes = export.to_excel(d, ["ticker", "company", "sector", "industry"], [],
                                  dt.datetime.now(), "title")
    wb = load_workbook(io.BytesIO(excel_bytes))
    ws = wb["Results"]
    # Formula prefixes and illegal char stripping
    assert ws.cell(row=2, column=1).value == "'+AAPL"
    assert ws.cell(row=2, column=2).value == "BadCompany"
    assert ws.cell(row=2, column=3).value == "'-Tech"
    assert ws.cell(row=2, column=4).value == "'@Software"

    csv_bytes = export.to_csv(d, ["ticker", "range52"])
    csv_str = csv_bytes.decode("utf-8-sig")
    assert "'+AAPL" in csv_str
    assert "49.0" in csv_str or "49" in csv_str


def test_doji_not_double_counted_with_dragonfly_or_gravestone():
    raw = frame_correctness(price=[10.0, 10.0, 10.0],
                            cdl_doji=[1.0, 1.0, 1.0],
                            cdl_doji_dragonfly=[0.0, 1.0, 0.0],
                            cdl_doji_gravestone=[0.0, 0.0, 1.0])
    d = add_derived(raw, today=dt.date(2026, 10, 20))
    # T0 has plain Doji
    assert "Doji" in d.loc[0, "signals"]
    # T1 has Dragonfly+, plain Doji must be suppressed
    assert "Dragonfly+" in d.loc[1, "signals"]
    assert "Doji" not in d.loc[1, "signals"].replace("Dragonfly+", "")
    # T2 has Gravestone-, plain Doji must be suppressed
    assert "Gravestone-" in d.loc[2, "signals"]
    assert "Doji" not in d.loc[2, "signals"].replace("Gravestone-", "")


def test_roe_with_negative_equity_excluded():
    # Net loss with negative equity gives positive ROE mathematically, but it's a loss
    d = frame_correctness(price=[10.0, 10.0], roe=[20.0, 20.0], debt_eq=[0.5, -2.0],
                          roic=[10.0, 10.0], gross_margin=[40.0, 40.0], oper_margin=[15.0, 15.0])
    # Filter "ROE > 0"
    pass_roe = apply_filters(d, {"roe": "Positive (>0%)"})
    assert pass_roe["ticker"].tolist() == ["T0"]

    # In scores, ROE with debt_eq < 0 is excluded
    s = add_scores(d)
    assert s.loc[0, "score_quality"] > s.loc[1, "score_quality"]


def test_minervini_stage2_full_template():
    p = P.BUILTIN["Stage 2 (Minervini)"]
    assert p.filters.get("sma200") == "SMA200 rising"
    assert p.filters.get("score_momentum") == "Over 70"
