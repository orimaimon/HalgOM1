"""Screener tests — run offline on synthetic data (no network)."""
import datetime as dt
import io
import json
import re
from pathlib import Path
from unittest.mock import patch

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from screener import export
from screener import groups as G
from screener import history
from screener import i18n as I
from screener import lifecycle
from screener import portfolio as PF
from screener import presets as P
from screener import signals as S
from screener.data.store import Store, load_universe
from screener.data.tradingview import clean, index_names, source_meta
from screener.derive import add_derived
from screener.fields import CANDLES, COLUMNS, GROUP_COLUMNS, INDEXES, VIEWS, provider_columns
from screener.data.tradingview import tv_fields, TV_MAPPING
from screener.filters import (ANY, FILTERS, FILTERS_BY_KEY, GROUPS, apply_filters, hide_junk, junk_mask,
                              option_labels, screen, search)
from screener.market import compute_regime, trend_strength
from screener.scores import add_scores

ROOT = Path(__file__).resolve().parent.parent

TODAY = dt.date(2026, 10, 18)


def make_raw(n: int = 300, seed: int = 7) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    data = {}
    for key, col in COLUMNS.items():
        if col.source != 'provider':
            continue
        if col.kind == "text":
            continue
        if col.kind == "date":
            data[key] = pd.Timestamp(TODAY) + pd.to_timedelta(rng.integers(-20, 90, n), unit="D")
        else:
            data[key] = rng.normal(10, 30, n)
    df = pd.DataFrame(data)
    df["ticker"] = [f"T{i:03d}" for i in range(n)]
    df["company"] = [f"Company {i}" for i in range(n)]
    df["exchange"] = rng.choice(["NASDAQ", "NYSE", "AMEX"], n)
    df["sector"] = rng.choice(["Technology Services", "Finance", "Health Technology"], n)
    df["industry"] = rng.choice(["Software", "Banks", "Biotech", "Semis"], n)
    df["country"] = rng.choice(["United States", "Israel", "Canada"], n, p=[0.8, 0.1, 0.1])
    df["sec_type"] = "stock"
    df["tv_symbol"] = df["exchange"] + ":" + df["ticker"]
    df["price"] = rng.uniform(1, 400, n)
    df["market_cap"] = 10 ** rng.uniform(7, 12.5, n)
    for c in ("volume", "avg_volume", "float_shares", "shares_out"):
        df[c] = 10 ** rng.uniform(4, 8, n)
    df["high52"] = df["price"] * rng.uniform(1.0, 1.8, n)
    df["low52"] = df["price"] * rng.uniform(0.4, 1.0, n)
    for c in ("sma20", "sma50", "sma150", "sma200", "sma20_prev", "sma50_prev", "sma200_prev", "close_prev",
              "sma50_w", "high_1m_prev", "high_3m_prev", "high_6m_prev", "high52_prev",
              "low_1m_prev", "low_3m_prev", "low52_prev"):
        df[c] = df["price"] * rng.uniform(0.7, 1.3, n)
    df["close_w"] = df["price"]
    df["high_1m"] = df["price"] * rng.uniform(1.0, 1.3, n)
    df["low_1m"] = df["price"] * rng.uniform(0.7, 1.0, n)
    df["high_5d"] = df["price"] * rng.uniform(1.0, 1.1, n)
    df["low_5d"] = df["price"] * rng.uniform(0.9, 1.0, n)
    df["atr_pct"] = rng.uniform(1, 8, n)
    df["rel_volume"] = rng.uniform(0.2, 6, n)
    df["rsi"] = rng.uniform(5, 95, n)
    df["rsi_w"] = rng.uniform(5, 95, n)
    for k in CANDLES:
        df[k] = (rng.random(n) < 0.05).astype(float)
    df["indexes"] = rng.choice(np.array(["S&P 500", "S&P 500, Nasdaq 100", "Russell 2000", None], dtype=object), n)
    df.loc[::17, "pe"] = np.nan          # missing fundamentals are common
    df.loc[::5, "pe"] = -abs(df.loc[::5, "pe"])
    return df


@pytest.fixture(scope="module")
def df():
    return add_derived(make_raw(), today=TODAY)


# ── registry ────────────────────────────────────────────────────────────────

def test_tv_fields_unique_and_views_valid():
    assert len(tv_fields()) > 50
    for name, cols in VIEWS.items():
        assert cols[0] == "ticker", name
        assert all(c in COLUMNS for c in cols), name


def test_filters_reference_known_columns():
    for f in FILTERS:
        assert f.column in COLUMNS, f.key
        assert f.tooltip, f"{f.key} has no tooltip"
        labels = [o.label for o in f.options]
        assert len(labels) == len(set(labels)), f"duplicate option label in {f.key}"


# ── filters ─────────────────────────────────────────────────────────────────

def test_every_option_returns_valid_mask(df):
    for f in FILTERS:
        for label in option_labels(f, df)[1:]:
            out = apply_filters(df, {f.key: label})
            assert len(out) <= len(df), (f.key, label)


def test_any_and_unknown_selections_are_ignored(df):
    assert len(apply_filters(df, {"pe": ANY, "price": ""})) == len(df)
    assert len(apply_filters(df, {"nope": "Over 5", "pe": "Not an option"})) == len(df)


def test_numeric_semantics_and_nan_never_passes(df):
    out = apply_filters(df, {"price": "Over $10", "pe": "Under 15"})
    assert (out["price"] > 10).all()
    assert out["pe"].between(0, 15, inclusive="neither").all()   # negative P/E is not "cheap"
    assert out["pe"].notna().all()


def test_combined_column_option(df):
    out = apply_filters(df, {"sma50": "SMA50 above SMA200"})
    assert (out["sma50"] > out["sma200"]).all()


def test_dynamic_country_options(df):
    labels = option_labels(FILTERS_BY_KEY["country"], df)
    assert labels[:3] == [ANY, "USA", "Foreign (ex-USA)"]
    assert (apply_filters(df, {"country": "USA"})["country"] == "United States").all()
    assert (apply_filters(df, {"country": "Foreign (ex-USA)"})["country"] != "United States").all()


def test_search(df):
    assert search(df, "T001, T002")["ticker"].tolist() == ["T001", "T002"]
    assert search(df, "AAA BBB").empty                       # all-caps words = ticker list
    assert len(search(df, "company 1")) >= 1                 # free text = substring
    assert len(search(df, "")) == len(df)


# ── presets ─────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("name", list(P.BUILTIN))
def test_builtin_presets_use_valid_options(name, df):
    p = P.BUILTIN[name]
    assert p.view in VIEWS
    assert p.sort is None or p.sort in COLUMNS
    for key, label in p.filters.items():
        f = FILTERS_BY_KEY[key]
        assert f.dynamic or label in option_labels(f), (name, key, label)


def test_user_presets_roundtrip(tmp_path):
    path = tmp_path / "presets.json"
    P.save_user("מסך שלי", P.Preset({"pe": "Under 15"}, view="Valuation"), path)
    assert P.load_user(path)["מסך שלי"].filters == {"pe": "Under 15"}
    P.delete_user("מסך שלי", path)
    assert P.load_user(path) == {}


# ── derive / scores ─────────────────────────────────────────────────────────

def test_adj_price_column(df):
    assert df["adj_price"].equals(df["price"])
    assert COLUMNS["adj_price"].kind == "price" and "Adj Price" == COLUMNS["adj_price"].label
    assert "adj_price" in VIEWS["Overview"]


def test_derived_columns(df):
    row = df.iloc[0]
    assert row["sma50_pct"] == pytest.approx((row["price"] / row["sma50"] - 1) * 100)
    assert df["range52"].dropna().between(0, 1).all()
    assert (df["high52_pct"].dropna() <= 0).all()
    assert (df["days_to_earnings"].dropna() >= 0).all()
    assert df.loc[df["earnings_date"] < pd.Timestamp(TODAY), "days_to_earnings"].isna().all()


def test_scores_rank_negative_pe_worst():
    d = pd.DataFrame({"pe": [5, 10, 50, -3], "ps": [1, 2, 5, 1], "pb": [1, 2, 5, 1],
                      "ev_ebitda": [5, 8, 30, -2], "pfcf": [5, 10, 40, -1]})
    s = add_scores(d)
    assert s["score_value"].iloc[0] == s["score_value"].max()
    assert s["score_value"].iloc[3] < s["score_value"].iloc[1]


def test_scores_in_range(df):
    for c in ("score", "score_value", "score_quality", "score_momentum"):
        assert df[c].dropna().between(0, 100).all()


# ── market regime (same rules as the HalgOM backtester) ──────────────────────

def _market(vix, spy, spy200, hyg, hyg20):
    return pd.DataFrame({"close": [spy, 1, hyg, vix], "change": 0.0, "sma20": [1, 1, hyg20, 1],
                         "sma50": 1.0, "sma200": [spy200, 1, 1, 1]}, index=["SPY", "QQQ", "HYG", "VIX"])


@pytest.mark.parametrize("args,expected", [
    ((15, 110, 100, 80, 79), "bull"),
    ((30, 90, 100, 80, 79), "bear"),
    ((20, 110, 100, 78, 79), "sideways"),
    ((17, 90, 100, 78, 79), "bear"),
    ((float("nan"), 110, 100, 80, 79), "unknown"),
])
def test_regime(args, expected):
    assert compute_regime(_market(*args)).label == expected


# ── cleaning ────────────────────────────────────────────────────────────────

def test_clean_drops_junk():
    raw = make_raw(5)
    raw["is_primary"] = [True, True, False, True, True]
    raw["currency"] = ["USD", "USD", "USD", "CAD", "USD"]
    raw["typespecs"] = [["common"]] * 5
    raw.loc[4, "ticker"] = raw.loc[0, "ticker"]   # duplicate ticker
    raw.loc[4, "market_cap"] = raw.loc[0, "market_cap"] * 2
    raw["earnings_date"] = 1795608000
    out = clean(raw)
    assert len(out) == 2
    assert out["ticker"].is_unique
    assert "is_primary" not in out.columns
    assert pd.api.types.is_datetime64_any_dtype(out["earnings_date"])


# ── store / offline fallback ────────────────────────────────────────────────

def test_store_cache_then_stale_fallback(tmp_path):
    store = Store(tmp_path)
    raw = make_raw(20)
    snap = load_universe(store, lambda: raw, dt.timedelta(minutes=15))
    assert snap.source == "live" and len(store.snapshot_dates()) == 1

    def boom():
        raise RuntimeError("offline")

    assert load_universe(store, boom, dt.timedelta(minutes=15)).source == "cache"
    stale = load_universe(store, boom, dt.timedelta(minutes=15), force=True)
    assert stale.source == "stale" and "offline" in stale.error and len(stale.df) == 20


def test_store_raises_without_cache(tmp_path):
    def boom():
        raise RuntimeError("offline")

    with pytest.raises(RuntimeError):
        load_universe(Store(tmp_path), boom, dt.timedelta(minutes=15))


# ── export ──────────────────────────────────────────────────────────────────

def test_excel_export(df):
    cols = VIEWS["Valuation"]
    data = export.to_excel(df.head(10), cols, [("P/E", "Under 15")], dt.datetime(2026, 10, 18, 15, 30), "t")
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Results"]
    headers = [c.value for c in ws[1]]
    assert headers[: len(cols)] == [COLUMNS[c].label for c in cols]
    assert headers[-3:] == export.JOURNAL_COLUMNS
    mcap = ws.cell(row=2, column=cols.index("market_cap") + 1)
    assert isinstance(mcap.value, float) and ',,,"B"' in mcap.number_format   # numeric, shown as B
    assert ws.cell(row=2, column=1).hyperlink.target.startswith("https://www.tradingview.com/symbols/")
    assert wb["Filters"]["A8"].value == "P/E"


def test_excel_billion_format_scales_correctly():
    # each trailing comma divides by 1,000 → billions need three
    assert export.EXCEL_FORMATS["money"].startswith('[>=1000000000]$0.00,,,"B"')


def test_watchlist_and_csv(df):
    wl = export.to_watchlist(df.head(12))
    assert wl.startswith("###")
    assert all(p.startswith("###") or ":" in p for p in wl.split(","))
    assert export.to_csv(df.head(3), ["ticker", "price"]).startswith(b"\xef\xbb\xbfTicker,Price")



# ══════════════════════════ Phase 2 ═════════════════════════════════════════

def frame(**cols) -> pd.DataFrame:
    """Minimal raw frame for targeted tests: every provider column present (NaN unless given)."""
    n = len(next(iter(cols.values())))
    base = {k: [np.nan] * n for k, c in COLUMNS.items() if c.source == 'provider' and c.kind != 'text'}
    base.update({"ticker": [f"X{i}" for i in range(n)], "company": [f"Co {i}" for i in range(n)],
                 "sector": ["Tech"] * n, "industry": ["Software"] * n, "country": ["United States"] * n,
                 "exchange": ["NASDAQ"] * n, "indexes": [None] * n,
                 "earnings_date": [pd.NaT] * n, "tv_symbol": [f"NASDAQ:X{i}" for i in range(n)]})
    for k in CANDLES:
        base[k] = [0.0] * n
    base.update(cols)
    return add_derived(pd.DataFrame(base), today=TODAY)


# ── registry ────────────────────────────────────────────────────────────────

def test_new_columns_registered_with_hebrew_help():
    hebrew = lambda t: any("֐" <= ch <= "׿" for ch in t)   # noqa: E731
    for k in [*CANDLES, "indexes", "sma150", "rsi_w", "sma50_w", "close_w", "sma50w_pct", "signals", "n_signals",
              "new", "range_1m_pct", "high52_prev"]:
        assert k in COLUMNS and hebrew(COLUMNS[k].help), k
    assert all(COLUMNS[k].kind == "flag" and next(tk for tk, v in TV_MAPPING.items() if v == k).startswith("Candle.") for k in CANDLES)
    assert {tk for tk, v in TV_MAPPING.items() if v in ("rsi_w", "sma50_w", "close_w")} == {"RSI|1W", "SMA50|1W", "close|1W"}
    assert all(hebrew(c.help) for c in GROUP_COLUMNS.values())
    assert "Signals" in GROUPS and "Signals" in VIEWS


def test_every_filter_group_is_known():
    assert {f.group for f in FILTERS} <= set(GROUPS)
    assert {"candle", "breakout", "consolidation", "vol_spike", "ma_cross", "ma_stack", "n_signals", "index",
            "sma150", "rsi_w", "sma50_w"} <= set(FILTERS_BY_KEY)


# ── signals ─────────────────────────────────────────────────────────────────

def test_signal_tags_have_hebrew_translations():
    import re
    # Every base tag must exist in SIGNALS_HE and contain no stray English letters except SMA and Cross
    for group in (S.CANDLE_SIGNALS, S.BREAKOUTS, S.BREAKDOWNS, S.CONSOLIDATION, S.CROSSES):
        for s in group:
            hebrew = I.SIGNALS_HE.get(s.tag)
            assert hebrew is not None, f"Missing translation for {s.tag}"
            if s.tag not in ("Golden Cross", "Death Cross"):
                # strip allowed English words and numbers, and punctuation
                stripped = re.sub(r"(SMA\d+|[0-9.x]+|\s|\+|-|↑|↓|')", "", hebrew)
                # the rest should be Hebrew (or empty)
                assert all("\u0590" <= ch <= "\u05FF" for ch in stripped), f"Stray English in {hebrew}"

    assert I.signal_tags("Engulfing+, Vol 2.8x, BO 1M", "he") == "בליעה+, מחזור ×2.8, פריצה חודשית"
    assert I.signal_tags("Engulfing+, Vol 2.8x, BO 1M", "en") == "Engulfing+, Vol 2.8x, BO 1M"
    assert I.signal_tags("Unknown, Vol 1.5x", "he") == "Unknown, מחזור ×1.5"
    # Hebrew cells are wrapped for the LTR grid; empty cells and non-strings stay as they are
    assert I.rtl_cell("בליעה+, SMA20↑SMA50") == "\u202bבליעה+, SMA20↑SMA50\u202c"
    assert I.rtl_cell("") == "" and I.rtl_cell(None) is None


def test_signals_column_combines_and_counts():
    d = frame(price=[105.0, 50.0, 10.0], rel_volume=[2.8, 1.0, np.nan],
              high_1m_prev=[100.0, 60.0, np.nan], high_3m_prev=[110.0, 70.0, np.nan],
              high_6m_prev=[110.0, 45.0, np.nan], high52_prev=[110.0, 49.0, np.nan],
              cdl_engulfing_bull=[1.0, 0.0, 0.0])
    assert d.loc[0, "signals"] == "Engulfing+, Vol 2.8x, BO 1M"
    assert d.loc[0, "n_signals"] == 3
    assert d.loc[1, "signals"] == "BO 52W"            # strongest breakout only, not 52W + 6M + 3M + 1M
    assert d.loc[2, "signals"] == "" and d.loc[2, "n_signals"] == 0   # missing data never signals

    from screener.app import display_values
    with patch("screener.i18n.HE", "he"), patch("screener.app.L", lambda: "he"):
        disp = display_values(d)
        assert disp.loc[0, "signals"] == "בליעה+, מחזור ×2.8, פריצה חודשית"
        assert disp.loc[1, "signals"] == "פריצה שנתית"

    with patch("screener.i18n.HE", "he"), patch("screener.app.L", lambda: "en"):
        disp = display_values(d)
        assert disp.loc[0, "signals"] == "Engulfing+, Vol 2.8x, BO 1M"
        assert disp.loc[1, "signals"] == "BO 52W"


def test_crosses_today_use_previous_bar():
    d = frame(price=[100.0, 100.0, 100.0], sma50=[90.0, 90.0, 80.0], sma200=[89.0, 91.0, 85.0],
              sma50_prev=[88.0, 92.0, 80.0], sma200_prev=[89.0, 91.0, 85.0])
    golden = apply_filters(d, {"ma_cross": "Golden Cross today (SMA50 crossed above SMA200)"})
    death = apply_filters(d, {"ma_cross": "Death Cross today (SMA50 crossed below SMA200)"})
    assert golden["ticker"].tolist() == ["X0"] and death["ticker"].tolist() == ["X1"]
    assert "Golden Cross" in d.loc[0, "signals"] and "Death Cross" in d.loc[1, "signals"]
    assert apply_filters(d, {"ma_cross": "SMA50 above SMA200"})["ticker"].tolist() == ["X0"]


def test_signal_filters():
    d = frame(price=[105.0, 95.0, 50.0], change=[3.0, -2.0, 0.5], rel_volume=[3.5, 2.2, 0.8],
              high_1m=[106.0, 120.0, 51.0], low_1m=[100.0, 90.0, 48.0], atr_pct=[3.0, 2.0, 1.0],
              high_1m_prev=[100.0, 120.0, 60.0], low_1m_prev=[90.0, 96.0, 40.0],
              cdl_hammer=[1.0, 0.0, 0.0], cdl_shooting_star=[0.0, 1.0, 0.0])
    t = lambda sel: apply_filters(d, sel)["ticker"].tolist()   # noqa: E731
    assert t({"candle": "Any bullish pattern"}) == ["X0"]
    assert t({"candle": "Any bearish pattern"}) == ["X1"]
    assert t({"candle": "Hammer"}) == ["X0"]
    assert t({"breakout": "New 1-month high (breakout)"}) == ["X0"]
    assert t({"breakout": "New 1-month low (breakdown)"}) == ["X1"]
    assert t({"vol_spike": "Volume 2x average"}) == ["X0", "X1"]
    assert t({"vol_spike": "Up on 2x volume"}) == ["X0"]
    assert t({"vol_spike": "Down on 2x volume"}) == ["X1"]
    # X0: 1M range 6% < 3 x ATR 3% ; X2: 6.25% vs 3 x 1% -> not tight relative to its own volatility
    assert t({"consolidation": "Volatility contraction (1M range < 3x ATR)"}) == ["X0"]
    assert t({"consolidation": "Tight 1-month range (<10%)"}) == ["X0", "X2"]
    assert t({"n_signals": "2 or more"}) == ["X0", "X1"]


def test_ma_stack_and_stage2_preset():
    #            pass   SMA150<SMA200  too close to low  too far from high
    d = frame(price=[100.0, 100.0, 100.0, 100.0], sma50=[90.0] * 4, sma150=[80.0, 70.0, 80.0, 80.0],
              sma200=[70.0, 75.0, 70.0, 70.0], sma200_prev=[65.0] * 4,
              perf_q=[30.0, 10.0, 10.0, 10.0], perf_h=[40.0, 10.0, 10.0, 10.0], perf_y=[50.0, 10.0, 10.0, 10.0],
              low52=[60.0, 60.0, 80.0, 60.0], high52=[110.0, 110.0, 110.0, 140.0])
    stage2 = P.BUILTIN["Stage 2 (Minervini)"]
    assert screen(d, stage2.filters, junk=False)["ticker"].tolist() == ["X0"]
    assert apply_filters(d, {"ma_stack": "Price > SMA50 > SMA150 > SMA200 (Stage 2)"})["ticker"].tolist() == \
        ["X0", "X2", "X3"]


# ── index ───────────────────────────────────────────────────────────────────

def test_index_names_from_tradingview_field():
    raw = [{"name": "NASDAQ Composite", "proname": "NASDAQ:IXIC"}, {"name": "NASDAQ 100", "proname": "NASDAQ:NDX"},
           {"name": "S&P 500", "proname": "SP:SPX"}]
    assert index_names(raw) == "S&P 500, Nasdaq 100"       # registry order, major indexes only
    assert index_names([{"name": "NASDAQ Composite", "proname": "NASDAQ:IXIC"}]) == ""
    assert index_names(None) is None
    out = clean(make_raw(3).assign(indexes=[raw, None, [{"name": "Dow", "proname": "DJ:DJI"}]]))
    assert sorted(out["indexes"].fillna("-")) == ["-", "DJIA", "S&P 500, Nasdaq 100"]


def test_index_filter_matches_whole_names():
    d = frame(price=[10.0] * 4, indexes=["S&P 500, Nasdaq 100", "Russell 2000", None, "S&P 500"])
    t = lambda label: apply_filters(d, {"index": label})["ticker"].tolist()   # noqa: E731
    assert t("S&P 500") == ["X0", "X3"]
    assert t("Nasdaq 100") == ["X0"]
    assert t("Russell 1000") == []
    assert t("Not in S&P 500") == ["X1", "X2"]
    assert t("Any major index") == ["X0", "X1", "X3"]
    assert set(INDEXES.values()) <= set(option_labels(FILTERS_BY_KEY["index"]))


# ── weekly ──────────────────────────────────────────────────────────────────

def test_weekly_filters():
    d = frame(price=[100.0, 100.0], close_w=[100.0, 100.0], sma50_w=[80.0, 120.0], rsi_w=[65.0, 25.0])
    assert d["sma50w_pct"].round(2).tolist() == [25.0, -16.67]
    t = lambda sel: apply_filters(d, sel)["ticker"].tolist()   # noqa: E731
    assert t({"sma50_w": "Price above SMA50 (weekly)"}) == ["X0"]
    assert t({"sma50_w": "Price 20% above SMA50 (weekly)"}) == ["X0"]
    assert t({"sma50_w": "Price 10% below SMA50 (weekly)"}) == ["X1"]
    assert t({"rsi_w": "Oversold (30)"}) == ["X1"]
    assert t({"rsi_w": "Above 50 (uptrend)"}) == ["X0"]


# ── hide junk ───────────────────────────────────────────────────────────────

def test_hide_junk():
    d = frame(price=[0.5, 20.0, 20.0, 20.0, 20.0], avg_volume=[1e6, 10e3, 1e6, 1e6, np.nan],
              market_cap=[1e9, 1e9, 10e6, 1e9, np.nan])
    assert junk_mask(d).tolist() == [True, True, True, False, False]    # missing values are not junk
    assert hide_junk(d)["ticker"].tolist() == ["X3", "X4"]
    assert len(hide_junk(d, on=False)) == 5
    assert screen(d, {}, junk=True)["ticker"].tolist() == ["X3", "X4"]


# ── data source metadata (stored with every snapshot) ───────────────────────

def test_source_meta():
    assert source_meta(pd.Series(["delayed_streaming_900"] * 3)) == {
        "provider": "tradingview", "update_mode": "delayed", "delay_min": 15, "raw": "delayed_streaming_900"}
    assert source_meta(pd.Series(["streaming"]))["update_mode"] == "realtime"
    assert source_meta(pd.Series(["endofday"]))["update_mode"] == "endofday"
    # anything unrecognized must never be shown to the user as real time
    assert source_meta(pd.Series(["something_new"]))["update_mode"] == "unknown"
    assert source_meta(None)["update_mode"] == "unknown"


def test_store_saves_source_meta_with_each_snapshot(tmp_path):
    raw = make_raw(10)
    raw.attrs["source"] = {"provider": "tradingview", "update_mode": "delayed", "delay_min": 15}
    store = Store(tmp_path)
    snap = load_universe(store, lambda: raw, dt.timedelta(minutes=15))
    assert snap.meta["update_mode"] == "delayed"
    day = store.snapshot_dates()[0]
    assert store.snapshot_meta(day)["source"]["provider"] == "tradingview"
    assert load_universe(store, lambda: raw, dt.timedelta(minutes=15)).meta["delay_min"] == 15   # from cache
    old = Store(tmp_path / "old")
    old.snap_dir.mkdir(parents=True)
    assert old.snapshot_meta(day) == {}            # snapshots from before metadata existed


def test_cache_missing_new_columns_is_refetched(tmp_path):
    store = Store(tmp_path)
    old = make_raw(10).drop(columns=list(CANDLES))
    load_universe(store, lambda: old, dt.timedelta(minutes=15))
    calls = []

    def fetch():
        calls.append(1)
        return make_raw(10)

    snap = load_universe(store, fetch, dt.timedelta(minutes=15), required=list(CANDLES))
    assert calls and snap.source == "live" and set(CANDLES) <= set(snap.df.columns)


def test_add_derived_tolerates_old_snapshot():
    old = make_raw(20).drop(columns=[*CANDLES, "sma150", "rsi_w", "indexes", "high_1m_prev"])
    d = add_derived(old, today=TODAY)
    assert d["sma150"].isna().all() and d["indexes"].isna().all()
    for f in FILTERS:                    # every filter still works, missing fields simply don't pass
        for label in option_labels(f, d)[1:]:
            apply_filters(d, {f.key: label})


# ── new since yesterday ─────────────────────────────────────────────────────

def test_previous_snapshot_date(tmp_path):
    store = Store(tmp_path)
    store.snap_dir.mkdir(parents=True)
    for day in ("2026-10-15", "2026-10-16", "2026-10-19"):
        make_raw(3).to_parquet(store.snap_dir / f"{day}.parquet")
    assert store.previous_snapshot_date(dt.date(2026, 10, 19)) == dt.date(2026, 10, 16)   # Monday -> Friday
    assert store.previous_snapshot_date(dt.date(2026, 10, 15)) is None


def test_new_since_yesterday(tmp_path):
    store = Store(tmp_path)
    store.snap_dir.mkdir(parents=True)
    raw = make_raw(40).assign(price=lambda d: d["price"].clip(lower=6))
    yesterday = raw.copy()
    yesterday.loc[:9, "price"] = 3.0                         # T000-T009 were under $5 yesterday
    yesterday = yesterday[yesterday["ticker"] != "T039"]     # T039 is a new listing
    yesterday.to_parquet(store.snap_dir / "2026-10-16.parquet")

    mon = dt.date(2026, 10, 19)
    prev = history.load_previous(store, mon)
    assert prev.day == dt.date(2026, 10, 16) and prev.missing == []
    today = add_derived(raw, today=TODAY)
    sel = {"price": "Over $5"}
    res = screen(today, sel, junk=False)
    new = history.mark_new(res, history.previous_tickers(prev, sel, junk=False))
    expected = {f"T{i:03d}" for i in range(10)} | {"T039"}
    assert set(res.loc[new == history.NEW, "ticker"]) == expected
    assert history.mark_new(res, None).eq("").all()


def test_new_since_yesterday_flags_old_snapshot(tmp_path):
    store = Store(tmp_path)
    store.snap_dir.mkdir(parents=True)
    make_raw(5).drop(columns=list(CANDLES)).to_parquet(store.snap_dir / "2026-10-16.parquet")
    mon = dt.date(2026, 10, 19)
    prev = history.load_previous(store, mon)
    assert set(prev.missing) == set(CANDLES)


# ── groups / heatmap ────────────────────────────────────────────────────────

def test_group_performance():
    d = frame(price=[10.0, 10.0, 10.0], sma50=[5.0, 20.0, np.nan], market_cap=[1e9, 3e9, 4e9],
              perf_w=[10.0, -10.0, 4.0], change=[1.0, 3.0, 2.0], sector=["A", "A", "B"])
    eq = G.group_performance(d, "sector").set_index("group")
    assert eq.loc["A", "stocks"] == 2 and eq.loc["A", "market_cap"] == 4e9
    assert eq.loc["A", "perf_w"] == pytest.approx(0.0)
    assert eq.loc["A", "pct_above_sma50"] == pytest.approx(50.0)
    assert np.isnan(eq.loc["B", "pct_above_sma50"])         # no SMA50 -> unknown, not 0%
    cap = G.group_performance(d, "sector", weight="cap").set_index("group")
    assert cap.loc["A", "perf_w"] == pytest.approx((10 * 1 - 10 * 3) / 4)
    assert G.group_performance(d, "sector")["group"].tolist() == ["A", "B"]      # largest first
    assert set(G.group_performance(d, "industry").columns) <= set(GROUP_COLUMNS)


@pytest.mark.parametrize("weight", ["equal", "cap"])
def test_group_performance_empty_input(weight):
    d = frame(price=[10.0], market_cap=[1e9], perf_w=[1.0]).iloc[:0]
    out = G.group_performance(d, "sector", weight=weight)
    assert out.empty and set(out.columns) <= set(GROUP_COLUMNS)


def test_group_performance_cap_weight_skips_missing_values():
    d = frame(price=[10.0] * 3, market_cap=[1e9, 3e9, 5e9], perf_w=[10.0, -10.0, np.nan], sector=["A"] * 3)
    cap = G.group_performance(d, "sector", weight="cap").set_index("group")
    assert cap.loc["A", "perf_w"] == pytest.approx(-5.0)       # the stock without Perf W adds no weight


def test_heatmap_nodes_and_figure():
    d = frame(price=[10.0] * 4, market_cap=[1e9, 3e9, 4e9, np.nan], change=[1.0, -3.0, 2.0, 5.0],
              sector=["A", "A", "B", "B"], industry=["a1", "a2", "b1", "b1"])
    n = G.heatmap_data(d)
    assert n["id"].is_unique and set(n["parent"]) - {""} <= set(n["id"])
    assert len(n[n["parent"] == ""]) == 1 and "X3" not in set(n["label"])   # no market cap -> no tile
    sector_a = n.set_index("id").loc["A"]
    assert sector_a["value"] == 4e9 and sector_a["color"] == pytest.approx((1 - 9) / 4)   # cap-weighted
    fig = G.heatmap_figure(d)
    assert fig.data[0].type == "treemap" and len(fig.data[0].ids) == len(n)


def test_heatmap_ids_survive_slashes_in_names():
    # industry "X/Y" and ticker "Y" in industry "X" collided as "S/X/Y" with a "/" separator
    d = frame(price=[10.0] * 2, market_cap=[1e9, 2e9], change=[1.0, 2.0], sector=["S", "S"],
              industry=["X/Y", "X"], ticker=["AAA", "Y"])
    n = G.heatmap_data(d)
    assert n["id"].is_unique and set(n["parent"]) - {""} <= set(n["id"])
    assert set(n["label"]) >= {"X/Y", "X", "Y", "AAA"}


# ── market trend strength (SPY ADX) ─────────────────────────────────────────

@pytest.mark.parametrize("adx,plus,minus,label", [
    (32, 30, 10, "Trend strong ▲"), (32, 10, 30, "Trend strong ▼"), (22, 30, 10, "Trend building ▲"),
    (12, 10, 30, "Trend weak"), (float("nan"), 1, 1, "Trend n/a"),
])
def test_trend_strength(adx, plus, minus, label):
    assert trend_strength(adx, plus, minus).label == label


def test_regime_rules_unchanged_by_adx():
    m = _market(15, 110, 100, 80, 79).assign(adx=5.0)
    assert compute_regime(m).label == "bull"


# ── app smoke test (Streamlit AppTest, offline) ─────────────────────────────

def test_selection_index():
    from screener.app import selection_index
    df = pd.DataFrame({"ticker": ["A", "B", "C", "B"]})
    df.index = [10, 20, 30, 40]

    # Ticker present
    assert selection_index(df, "B") == {"selection": {"rows": [1]}}
    # Ticker absent
    assert selection_index(df, "D") is None
    # Empty dataframe
    assert selection_index(pd.DataFrame(columns=["ticker"]), "A") is None
    # No remembered ticker
    assert selection_index(df, None) is None
    assert selection_index(df, "") is None


def test_app_smoke(tmp_path, monkeypatch):
    import streamlit as st
    from streamlit.testing.v1 import AppTest

    import screener.data.store as store_mod
    import screener.data.tradingview as tv

    raw = clean(make_raw(300).assign(is_primary=True, currency="USD"))
    raw.attrs["source"] = {"provider": "tradingview", "update_mode": "delayed", "delay_min": 15}
    market = pd.DataFrame({"close": [500.0, 400, 80, 15], "change": 0.5, "sma20": [490.0, 390, 79, 16],
                           "sma50": 480.0, "sma200": [450.0, 350, 78, 17], "adx": [30.0, 20, 15, 10],
                           "di_plus": 25.0, "di_minus": 15.0}, index=["SPY", "QQQ", "HYG", "VIX"])
    monkeypatch.setattr(store_mod, "DATA_DIR", tmp_path)
    monkeypatch.setattr(P, "PRESETS_PATH", tmp_path / "presets.json")
    monkeypatch.setattr(PF, "PORTFOLIOS_PATH", tmp_path / "portfolios.json")
    monkeypatch.setattr(export, "EXPORT_DIR", tmp_path / "exports")
    monkeypatch.setattr(tv, "fetch_universe", lambda: raw)
    monkeypatch.setattr(tv, "fetch_market", lambda: market)
    stops = []
    monkeypatch.setattr(lifecycle, "start_idle_watchdog", lambda: None)
    monkeypatch.setattr(lifecycle, "stop_server", lambda reason: stops.append(reason))
    snaps = tmp_path / "snapshots"
    snaps.mkdir()
    prev_session = store_mod.session_date(dt.datetime.now() - dt.timedelta(days=4))
    raw.head(250).to_parquet(snaps / f"{prev_session}.parquet")
    st.cache_data.clear()
    st.cache_resource.clear()

    at = AppTest.from_file(str(ROOT / "screener" / "app.py"), default_timeout=60).run()
    assert not at.exception, at.exception
    assert at.session_state["hide_junk"] is True
    # Hebrew is the default; values stay English internally
    assert at.session_state["lang"] == "he"
    assert any("תוצאות" in m.value for m in at.markdown)
    assert any("בעיכוב 15 דק'" in c.value for c in at.caption)
    assert at.selectbox(key="f_pe").options[1] == "נמוך (<15)"          # displayed in Hebrew...
    at.selectbox(key="f_pe").select("Low (<15)").run()                   # ...stored in English
    assert at.session_state["f_pe"] == "Low (<15)"
    at.selectbox(key="f_pe").select("Any").run()
    # English (a separate AppTest: its selectbox bookkeeping does not follow a language switch mid-test;
    # the real switch is checked in the browser)
    en = AppTest.from_file(str(ROOT / "screener" / "app.py"), default_timeout=60)
    en.session_state["lang"] = "en"
    en.run()
    assert not en.exception, en.exception
    assert any("results" in m.value for m in en.markdown)
    assert any("delayed 15 min" in c.value for c in en.caption)
    assert en.selectbox(key="f_pe").options[1] == "Low (<15)"
    assert at.session_state["prev"] is not None                     # yesterday's snapshot was found

    at.selectbox(key="preset").select("Stage 2 (Minervini)").run()
    assert not at.exception, at.exception
    assert at.session_state["f_ma_stack"] == "Price > SMA50 > SMA150 > SMA200 (Stage 2)"
    at.toggle(key="hide_junk").set_value(False).run()
    at.toggle(key="only_new").set_value(True).run()
    at.session_state["view"] = "Signals"
    at.run()
    assert not at.exception, at.exception

    for page in ("Groups", "Guide", "Portfolios", "Heatmap"):
        at.session_state["page"] = page
        at.run()
        assert not at.exception, (page, at.exception)
    at.selectbox(key="h_universe").select("All stocks (top 1,000)").run()
    at.selectbox(key="h_color").select("perf_m").run()
    at.session_state["page"] = "Groups"
    at.run()
    at.session_state["g_by"] = "Industry"
    at.run()
    assert not at.exception, at.exception
    # cap-weighted averages over an empty screen (used to raise ValueError)
    at.session_state["g_weight"] = "Market cap"
    at.session_state["g_results"] = True
    at.session_state["query"] = "NO SUCH TICKER"
    at.run()
    assert not at.exception, at.exception
    at.session_state["query"] = ""
    at.session_state["g_results"] = False
    at.run()

    at.session_state["page"] = "Screener"
    at.run()
    assert not at.exception, at.exception
    # filters and switches survive a visit to the other pages
    assert at.session_state["f_ma_stack"] == "Price > SMA50 > SMA150 > SMA200 (Stage 2)"
    assert at.session_state["hide_junk"] is False
    assert json.loads((tmp_path / "cache" / "latest_meta.json").read_text())["source"]["update_mode"] == "delayed"

    # Exit: goodbye page, the server is asked to stop exactly once (even if the page reruns)
    at.button(key="exit").click().run()
    assert not at.exception, at.exception
    assert stops == ["Exit button"]
    at.run()
    assert stops == ["Exit button"]
    assert not at.dataframe                                      # only the goodbye page is left



# ── review fixes ────────────────────────────────────────────────────────────

def test_candle_registry_carries_its_own_tags():
    assert all(len(c) == 4 and c[2] for c in CANDLES.values())
    assert [s.tag for s in S.CANDLE_SIGNALS] == [c[2] + c[1] for c in CANDLES.values()]


def test_ma_comparisons_identical_across_filters():
    d = frame(price=[10.0] * 3, sma20=[3.0, 1.0, np.nan], sma50=[2.0, 2.0, 2.0], sma200=[1.0, 3.0, np.nan])
    for label in ("SMA50 above SMA200", "SMA50 below SMA200"):
        assert apply_filters(d, {"sma50": label}).equals(apply_filters(d, {"ma_cross": label}))
    for label in ("SMA20 above SMA50", "SMA20 below SMA50"):
        assert apply_filters(d, {"sma20": label}).equals(apply_filters(d, {"ma_cross": label}))


def test_six_month_breakdown():
    d = frame(price=[50.0, 50.0], low52_prev=[40.0, 40.0], low_6m_prev=[55.0, 45.0], low_3m_prev=[56.0, 48.0],
              low_1m_prev=[57.0, 49.0])
    assert apply_filters(d, {"breakout": "New 6-month low (breakdown)"})["ticker"].tolist() == ["X0"]
    assert d["signals"].tolist() == ["BD 6M", ""]                 # strongest breakdown only
    labels = option_labels(FILTERS_BY_KEY["breakout"])
    for months in ("1-month", "3-month", "6-month", "52-week"):
        assert f"New {months} high (breakout)" in labels and f"New {months} low (breakdown)" in labels



# ── lifecycle: Exit button + idle watchdog ──────────────────────────────────

def test_idle_watch_counts_consecutive_empty_polls():
    w = lifecycle.IdleWatch(idle_polls=3)
    assert [w.tick(n) for n in (0, 0, 1, 0, 0)] == [False] * 5     # an open tab resets the count
    assert w.tick(0) is True
    w = lifecycle.IdleWatch(idle_polls=2)
    assert [w.tick(n) for n in (0, None, 0, 0)] == [False, False, False, True]   # unknown never stops


def test_idle_watchdog_stops_server_once_idle():
    import threading
    counts = iter([2, 1, 0, 0, 0, 0, 0])
    stopped = threading.Event()
    reasons = []

    def stop(reason):
        reasons.append(reason)
        stopped.set()

    t = lifecycle.start_idle_watchdog(count=lambda: next(counts, 0), stop=stop, poll=0.01, idle_polls=3)
    assert stopped.wait(5) and reasons == ["no browser tab for 3 checks"]
    t.join(5)
    assert not t.is_alive()


def test_active_sessions_outside_a_server_is_unknown():
    assert lifecycle.active_sessions() is None


# ── launcher (run_screener.bat) ─────────────────────────────────────────────

BAT = (ROOT / "run_screener.bat").read_bytes()
# The update block exactly as in the first launcher users installed. git pull may rewrite the
# running .bat, and cmd resumes at the same byte offset — so this block must end at byte 785.
UPDATE_BLOCK = b"\r\n".join([
    b'if not defined SCREENER_NO_UPDATE if exist ".git" (',
    b"    where git >nul 2>nul",
    b"    if not errorlevel 1 (",
    b"        echo [2/3] Checking for updates...",
    b"        git pull --ff-only --quiet || echo        update skipped - running the current version",
    b"    )",
    b")",
    b"",
])


def test_launcher_is_ascii_crlf():
    BAT.decode("ascii")
    assert b"\n" not in BAT.replace(b"\r\n", b"")


def test_launcher_update_block_frozen_at_offset():
    assert BAT[529:785] == UPDATE_BLOCK
    assert BAT[785:].startswith(b"\r\nrem ")        # an old launcher resumes here, at a line boundary


def test_launcher_git_pull_is_switched_off_and_restored():
    text = BAT.decode("ascii")
    assert 'set "SCREENER_NO_UPDATE=1"' in BAT[:529].decode("ascii")    # the frozen git pull never runs
    after = BAT[785:].decode("ascii")
    assert after.index('set "SCREENER_NO_UPDATE=%SKIP_UPDATE%"') < after.index("screener.update")
    # an old launcher that just pulled this file resumes at byte 785: everything it needs is set below it
    assert 'set "UV=1"' in after and 'set "PY=' in after
    assert "where uv >nul 2>nul\r\nif errorlevel 1 goto :fail" not in text     # no uv: fall back to Python


def test_launcher_update_step_is_one_line_that_restarts():
    lines = [ln for ln in BAT.decode("ascii").split("\r\n") if "screener.update" in ln and not ln.startswith("rem")]
    assert len(lines) == 1                         # update.py may rewrite this file: cmd has read the line
    line = lines[0]
    assert line.startswith("if not defined SCREENER_NO_UPDATE if not defined SCREENER_UPDATED ")
    assert "(" not in line and ")" not in line     # %~f0 may contain ")" (a folder "HalgOM (1)")
    # `call "%~f0"` keeps the path quoted in this cmd and in the new one: a folder "Stocks & Co" works
    # (`cmd /c ""%~f0""` left it unquoted here, so "&" split the line)
    assert line.endswith('& if errorlevel 3 set "SCREENER_UPDATED=1" & cmd /c call "%~f0" & exit /b')
    from screener import update
    assert update.UPDATED == 3


def test_launcher_rebuilds_a_copied_venv():
    text = BAT.decode("ascii")
    probe = '"%PY%" -c "import sys; assert sys.version_info.minor >= 11"'   # starts, and is new enough
    assert probe in text and 'rmdir /s /q "screener\\.venv"' in text
    assert text.index(probe) < text.index("uv venv --python 3.12")


def test_launcher_failed_install_still_updates():
    """A broken install must not block the update that may fix it."""
    text = BAT.decode("ascii")
    update = text.index("-m screener.update")
    assert 'if errorlevel 1 set "INSTALL_FAILED=1"' in text
    assert text.index('set "INSTALL_FAILED="') < text.index("pip install") < update
    assert update < text.index("if defined INSTALL_FAILED goto :fail")


def test_launcher_explains_a_busy_port():
    text = BAT.decode("ascii")
    assert "if errorlevel 1 goto :portbusy" in text and "Get-NetTCPConnection -LocalPort 8501" in text
    assert "using port 8501" in text


def test_launcher_labels_and_port():
    import re
    text = BAT.decode("ascii")
    labels = set(re.findall(r"^:(\w+)", text, re.M))
    targets = set(re.findall(r"(?:goto|call) :(\w+)", text)) - {"eof"}
    assert targets <= labels, targets - labels
    assert "--server.port 8501" in text and "http://localhost:8501/_stcore/health" in text
    for line in text.splitlines():              # "rem-" (no space) is not a comment in cmd
        assert not re.match(r"\s*rem\S", line, re.I), line


def test_launcher_echo_inside_blocks_has_no_parentheses():
    depth = 0
    for line in BAT.decode("ascii").splitlines():
        stripped = line.strip()
        if depth and stripped.lower().startswith("echo"):
            assert "(" not in stripped and ")" not in stripped, line   # would close the block early
        if stripped.endswith("(") and not stripped.lower().startswith("rem"):
            depth += 1
        elif stripped == ")":
            depth -= 1
    assert depth == 0



# ── Hebrew display (i18n.py) ────────────────────────────────────────────────

ALLOWED_ENGLISH = re.compile(r"\b(SMA\d*|ATR|RSI|EPS|PEG|ROA|ROE|ROIC|FCF|EV|EBITDA|P|E|S|B|M|K|x|"
                             r"Float|Gap|Beta|Payout|Fwd|HalgOM|Value|Quality|Momentum|Golden|Death|Cross|"
                             r"Minervini|Nasdaq|DJIA|Russell|TradingView|Doji|Marubozu|Harami|Engulfing|Hammer|"
                             r"Inverted|Hanging|Man|Shooting|Star|Morning|Evening|Three|White|Soldiers|Black|Crows|"
                             r"Bullish|Bearish|Kicking|Abandoned|Baby|Dragonfly|Gravestone|OCF)\b")


def english_left(text: str) -> set[str]:
    text = re.sub(r"\d[\d.,]*[KMBx]", "", text)          # amounts like $200B, 50K, 3x
    return set(re.findall(r"[A-Za-z]+", ALLOWED_ENGLISH.sub("", text)))


def test_every_filter_and_option_has_hebrew():
    for f in FILTERS:
        label = I.filter_label(f.key, f.label, I.HE)
        assert f.key in I.FILTERS_HE and not english_left(label), (f.key, label)
        for o in f.options:
            he = I.option(o.label, I.HE)
            index_name = o.label in set(I.option(x, I.EN) for x in ("S&P 500", "Nasdaq 100", "DJIA",
                                                                       "Russell 1000", "Russell 2000"))
            assert index_name or he != o.label, (f.key, o.label)
            assert not english_left(he), (f.key, o.label, he)


def test_columns_views_presets_groups_have_hebrew():
    for k, c in COLUMNS.items():
        assert not english_left(I.column(k, c.label, I.HE)), k
    for k, c in GROUP_COLUMNS.items():
        assert not english_left(I.group_column(k, c.label, I.HE)), k
    for name in [*VIEWS, *GROUPS, "Screener", "Groups", "Heatmap", "Portfolios"]:
        assert I.ui(name, I.HE) != name, name
    for name in P.BUILTIN:
        assert I.preset(name, I.HE) != name, name


def test_english_mode_is_unchanged():
    assert I.option("Under 15", I.EN) == "Under 15"
    assert I.column("pe", "P/E", I.EN) == "P/E"
    assert I.value("sector", "Finance", I.EN) == "Finance"
    assert I.ui("Refresh data", I.EN) == "Refresh data"


def test_hebrew_values_and_patterns():
    assert I.value("sector", "Finance", I.HE) == "פיננסים"
    assert I.value("country", "Israel", I.HE) == "ישראל"
    assert I.value("industry", "Semiconductors", I.HE) == "Semiconductors"     # industries stay English
    assert I.option("Under -10%", I.HE) == "מתחת ל -10%"
    assert I.option("$5 to $10", I.HE) == "$5 עד $10"
    assert I.option("Price 10% below SMA50 (weekly)", I.HE) == "מחיר 10% מתחת ל-SMA50 (שבועי)"
    assert I.option("Something new", I.HE) == "Something new"                   # unknown -> English fallback


def test_excel_export_in_hebrew(df):
    cols = ["ticker", "sector", "pe"]
    data = export.to_excel(df.head(3).assign(sector="פיננסים"), cols, [("מכפיל רווח (P/E)", "נמוך (<15)")],
                           dt.datetime(2026, 10, 18, 15, 30), "ערך", headers={"ticker": "סימול", "pe": "P/E"},
                           t=lambda x: I.ui(x, I.HE), rtl=True)
    wb = load_workbook(io.BytesIO(data))
    ws = wb["תוצאות"]
    assert ws.sheet_view.rightToLeft
    assert [c.value for c in ws[1]][:3] == ["סימול", "Sector", "P/E"]          # missing header -> English label
    assert [c.value for c in ws[1]][-3:] == ["סטטוס / פעולה", "מחיר כניסה", "הערות"]
    assert ws.cell(row=2, column=2).value == "פיננסים"
    assert wb["פילטרים"]["A8"].value == "מכפיל רווח (P/E)"

# ── user guide (screener/guide.html) ────────────────────────────────────────

GUIDE = ROOT / "screener" / "guide.html"


def test_guide_covers_presets_and_ui_names():
    """The guide quotes the Hebrew UI; renaming a preset, page, tab or view must update the guide too."""
    text = GUIDE.read_text(encoding="utf-8")
    assert "<script" not in text.lower()
    for name in I.PRESETS_HE.values():
        assert f"<td>{name}</td>" in text, name
    quoted = set(re.findall(r'<span class="ui">([^<]+)</span>', text))
    for en in ["Screener", "Groups", "Heatmap", "Portfolios", *GROUPS, *VIEWS]:
        assert I.ui(en, I.HE) in quoted, en
    known = (set(I.UI.values()) | set(I.FILTERS_HE.values()) | set(I.OPTIONS_HE.values())
             | set(I.LANGS.values()) | {"HalgOM Screener", "Excel", "CSV"})
    assert quoted <= known, quoted - known
    for anchor in re.findall(r'href="#([\w-]+)"', text):
        assert f'id="{anchor}"' in text, anchor
