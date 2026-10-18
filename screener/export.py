"""
Exports: Excel (formatted), CSV, TradingView watchlist, plain ticker list.

Excel keeps every number numeric — display formats (2.3B, 5.20%) are Excel number formats,
so sorting and formulas keep working.
"""
from __future__ import annotations

import datetime as dt
import io
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from screener.data.store import DATA_DIR
from screener.fields import COLUMNS

EXPORT_DIR = DATA_DIR / "exports"

EXCEL_FORMATS = {
    "price": "#,##0.00",
    "pct": '0.00"%"',
    "ratio": "0.00",
    "int": "#,##0",
    "money": '[>=1000000000]$0.00,,,"B";[>=1000000]$0.0,,"M";$#,##0',
    "count": '[>=1000000000]0.00,,,"B";[>=1000000]0.00,,"M";#,##0',
    "date": "yyyy-mm-dd",
    "days": "0",
    "score": "0",
    "range": "0%",
    "flag": "0",
}
JOURNAL_COLUMNS = ["Status / Action", "Target Entry", "Notes"]
_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_JOURNAL_FILL = PatternFill("solid", fgColor="FFF4D6")
_GREEN, _RED = Font(color="1A7F37"), Font(color="CF222E")
_SIGNED = {"change", "change_abs", "gap", "premarket", "postmarket", "perf_w", "perf_m", "perf_q", "perf_h",
           "perf_ytd", "perf_y", "perf_5y", "eps_growth", "eps_growth_q", "sales_growth", "sales_growth_q",
           "target_upside", "sma20_pct", "sma50_pct", "sma150_pct", "sma200_pct", "sma50w_pct"}


def tradingview_url(tv_symbol: str) -> str:
    return f"https://www.tradingview.com/symbols/{tv_symbol.replace(':', '-')}/"


def to_excel(df: pd.DataFrame, columns: list[str], filters: list[tuple[str, str]],
             fetched_at: dt.datetime, title: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    cols = [c for c in columns if c in df.columns]

    headers = [COLUMNS[c].label for c in cols] + JOURNAL_COLUMNS
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for j in range(len(cols) + 1, len(headers) + 1):
        ws.cell(row=1, column=j).fill = PatternFill("solid", fgColor="8A6D1F")

    for row in df[cols].itertuples(index=False):
        ws.append([_cell_value(v) for v in row])

    n = len(df)
    symbols = df["tv_symbol"].tolist() if "tv_symbol" in df.columns else [None] * n
    for j, key in enumerate(cols, start=1):
        kind = COLUMNS[key].kind
        letter = get_column_letter(j)
        fmt = EXCEL_FORMATS.get(kind)
        for i in range(2, n + 2):
            cell = ws.cell(row=i, column=j)
            if fmt:
                cell.number_format = fmt
            if key == "ticker" and symbols[i - 2]:
                cell.hyperlink = tradingview_url(symbols[i - 2])
                cell.font = Font(color="0B57D0", underline="single", bold=True)
        if key in _SIGNED and n:
            rng = f"{letter}2:{letter}{n + 1}"
            ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=_GREEN))
            ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=_RED))
        width = (32 if key in ("company", "signals") else 22 if key in ("industry", "sector", "indexes")
                 else max(10, len(COLUMNS[key].label) + 3))
        ws.column_dimensions[letter].width = width
    for j in range(len(cols) + 1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 18 if headers[j - 1] != "Notes" else 40
        for i in range(2, n + 2):
            ws.cell(row=i, column=j).fill = _JOURNAL_FILL

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{n + 1}"
    ws.row_dimensions[1].height = 30

    meta = wb.create_sheet("Filters")
    meta.append(["HalgOM Screener export"])
    meta["A1"].font = Font(bold=True, size=14)
    meta.append(["Screen", title or "(custom)"])
    meta.append(["Data as of", fetched_at.strftime("%Y-%m-%d %H:%M")])
    meta.append(["Exported", dt.datetime.now().strftime("%Y-%m-%d %H:%M")])
    meta.append(["Results", n])
    meta.append([])
    meta.append(["Filter", "Value"])
    for c in meta[7]:
        c.font = Font(bold=True)
    for label, value in filters or [("(none)", "")]:
        meta.append([label, value])
    meta.append([])
    meta.append(["Source: TradingView scanner (delayed ~15 min). Not investment advice."])
    meta.column_dimensions["A"].width = 28
    meta.column_dimensions["B"].width = 36

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_value(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if hasattr(v, "item"):  # numpy scalar
        return v.item()
    return v


def to_csv(df: pd.DataFrame, columns: list[str]) -> bytes:
    cols = [c for c in columns if c in df.columns]
    out = df[cols].rename(columns={c: COLUMNS[c].label for c in cols})
    return out.to_csv(index=False).encode("utf-8-sig")  # BOM so Excel opens it correctly


def to_watchlist(df: pd.DataFrame, group_by: str | None = "sector") -> str:
    """TradingView watchlist import file: "###Section,EXCH:SYM,..." (sections optional)."""
    if group_by and group_by in df.columns:
        parts = []
        for name, g in df.groupby(df[group_by].fillna("Other"), sort=True):
            parts.append(f"###{_section(name)}")
            parts.extend(g["tv_symbol"].tolist())
        return ",".join(parts)
    return ",".join(df["tv_symbol"].tolist())


def _section(name: str) -> str:
    return re.sub(r"[,#]", " ", str(name)).strip() or "Other"


def tickers_text(df: pd.DataFrame) -> str:
    return ", ".join(df["ticker"].tolist())


def file_stem(title: str, when: dt.datetime | None = None) -> str:
    when = when or dt.datetime.now()
    slug = re.sub(r"[^A-Za-z0-9]+", "-", title or "screen").strip("-").lower() or "screen"
    return f"{when:%Y-%m-%d_%H%M}_{slug}"


def save_copy(data: bytes | str, filename: str, folder: Path | None = None) -> Path:
    folder = folder or EXPORT_DIR
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / filename
    if isinstance(data, str):
        path.write_text(data, encoding="utf-8")
    else:
        path.write_bytes(data)
    return path
